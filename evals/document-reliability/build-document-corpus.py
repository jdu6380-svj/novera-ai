#!/usr/bin/env python3
"""Build real local document fixtures from corpus-spec.v1.json."""

from __future__ import annotations

import hashlib
import json
import mimetypes
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parent
SPEC_PATH = ROOT / "corpus-spec.v1.json"
FIXTURE_DIR = ROOT / "fixtures"
MANIFEST_PATH = ROOT / "manifest.v1.json"


def load_cjk_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path("/System/Library/Fonts/PingFang.ttc"),
        Path("/System/Library/Fonts/STHeiti Light.ttc"),
        Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf"),
    ]
    for candidate in candidates:
        if candidate.is_file():
            try:
                return ImageFont.truetype(str(candidate), size=size, index=0)
            except OSError:
                continue
    return ImageFont.load_default()


def render_scan_page(text: str, destination: Path) -> None:
    image = Image.new("RGB", (1275, 1650), "white")
    draw = ImageDraw.Draw(image)
    font = load_cjk_font(42)
    y = 180
    for line in text.splitlines():
        draw.text((110, y), line, fill="#222222", font=font)
        y += 92
    image.save(destination, format="PNG")


def build_pdf(document: dict, path: Path) -> None:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    pages = document["content"]["pages"]
    page_images: list[Path] = []
    for index, page in enumerate(pages):
        if isinstance(page, dict) and page.get("kind") == "image":
            image_path = FIXTURE_DIR / f".{document['document_id']}-scan-{index + 1}.png"
            render_scan_page(page["text"], image_path)
            page_images.append(image_path)

    pdf = canvas.Canvas(str(path), pagesize=letter, pageCompression=0)
    page_width, page_height = letter
    pdf.setTitle(document["document_title"])
    for index, page in enumerate(pages):
        if isinstance(page, dict) and page.get("kind") == "image":
            image_path = FIXTURE_DIR / f".{document['document_id']}-scan-{index + 1}.png"
            pdf.drawImage(str(image_path), 0, 0, width=page_width, height=page_height)
        else:
            text = page["text"] if isinstance(page, dict) else page
            pdf.setFont("STSong-Light", 14)
            y = page_height - 64
            for line in text.splitlines():
                if line:
                    pdf.drawString(58, y, line)
                y -= 26
                if y < 42:
                    pdf.showPage()
                    pdf.setFont("STSong-Light", 14)
                    y = page_height - 64
        pdf.showPage()
    pdf.save()
    for image_path in page_images:
        image_path.unlink()


def build_docx(document: dict, path: Path) -> None:
    content = document["content"]
    paragraphs = content.get("paragraphs", [])
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)

    if paragraphs:
        title = doc.add_paragraph(style="Title")
        title.add_run(paragraphs[0])
        title.paragraph_format.space_after = Pt(14)
        for paragraph in paragraphs[1:]:
            if paragraph[:1].isdigit() and ". " in paragraph[:4]:
                doc.add_paragraph(paragraph, style="List Number")
            else:
                doc.add_paragraph(paragraph)

    table_rows = content.get("table")
    if table_rows:
        table = doc.add_table(rows=0, cols=len(table_rows[0]))
        table.style = "Table Grid"
        for row in table_rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row, strict=True):
                cell.text = str(value)

    doc.core_properties.title = document["document_title"]
    doc.save(path)


def build_png(document: dict, path: Path) -> None:
    content = document["content"]
    image = Image.new("RGB", (1200, 800), "white")
    draw = ImageDraw.Draw(image)
    title_font = load_cjk_font(48)
    body_font = load_cjk_font(34)
    draw.text((90, 70), content["title"], fill="#1D2939", font=title_font)
    draw.line((90, 145, 1110, 145), fill="#98A2B3", width=2)
    values = [line.rsplit("：", 1)[-1] for line in content["lines"][:2]]
    labels = [line.split("：", 1)[0] for line in content["lines"][:2]]
    bars = [520, 760]
    y = 235
    for label, value, width in zip(labels, values, bars, strict=True):
        draw.text((90, y), label, fill="#344054", font=body_font)
        draw.rounded_rectangle((90, y + 66, 90 + width, y + 118), radius=8, fill="#2F6B5B")
        draw.text((105 + width, y + 68), value, fill="#101828", font=body_font)
        y += 220
    draw.text((90, 730), content["lines"][-1], fill="#667085", font=load_cjk_font(24))
    image.save(path, format="PNG")


def create_fixture(document: dict) -> bytes:
    path = FIXTURE_DIR / document["file_name"]
    if path.exists():
        raise FileExistsError(f"Refusing to overwrite existing fixture: {path}")

    file_format = document["file_format"]
    content = document["content"]
    if file_format in {"txt", "markdown"}:
        text = content["text"]
        if document["document_id"] == "doc-014":
            path.write_bytes(text.encode("utf-8"))
        else:
            path.write_text(text, encoding="utf-8", newline="")
    elif file_format == "docx":
        build_docx(document, path)
    elif file_format == "pdf":
        build_pdf(document, path)
    elif file_format == "png":
        build_png(document, path)
    elif file_format in {"xlsx", "pptx"}:
        return b""
    else:
        raise ValueError(f"Unsupported fixture builder format: {file_format}")
    return path.read_bytes()


def mime_type(document: dict) -> str:
    return mimetypes.guess_type(document["file_name"])[0] or "application/octet-stream"


def main() -> None:
    spec = json.loads(SPEC_PATH.read_text(encoding="utf-8"))
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    if MANIFEST_PATH.exists():
        raise FileExistsError(f"Refusing to overwrite existing manifest: {MANIFEST_PATH}")

    documents = []
    for source in spec["documents"]:
        if source["file_format"] in {"xlsx", "pptx"}:
            continue
        content_bytes = create_fixture(source)
        if not content_bytes:
            raise RuntimeError(f"Fixture was not written: {source['file_name']}")
        file_format = source["file_format"]
        documents.append(
            {
                "document_id": source["document_id"],
                "file_name": source["file_name"],
                "file_format": file_format,
                "source_type": source["source_type"],
                "company_name": source["company_name"],
                "document_title": source["document_title"],
                "document_version": source["document_version"],
                "mime_type": mime_type(source),
                "document_hash": hashlib.sha256(content_bytes).hexdigest(),
                "expected_capability": source["expected_capability"],
                "expected_behavior": source["expected_behavior"],
                "execution_status": "NOT_RUN",
                "actual_behavior": "Awaiting local parser evaluation.",
                "known_limitations": source["known_limitations"],
            }
        )

    manifest = {
        "schema_version": spec["schema_version"],
        "dataset_name": spec["dataset_name"],
        "source_type": spec["source_type"],
        "notice": spec["notice"],
        "documents": documents,
        "cases": [
            {
                "case_id": question["case_id"],
                "document_id": document["document_id"],
                "document_name": document["file_name"],
                "question": question["question"],
                "question_type": question["question_type"],
                "expected_answer": question["expected_answer"],
                "evidence_text": question["evidence"],
                "evidence_location": [item["evidence_location"] for item in question["evidence"]],
                "answerability": question["answerability"],
                "expected_behavior": question["expected_behavior"],
                "actual_result": "NOT_RUN",
                "verification_status": "NOT_RUN",
                **({"related_document_ids": question["related_document_ids"]} if question.get("related_document_ids") else {}),
            }
            for document in spec["documents"]
            for question in document["questions"]
        ],
    }
    MANIFEST_PATH.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Created {len(documents)} non-office fixtures and {len(manifest['cases'])} question cases.")
    print(f"Office fixtures are created by build-office-corpus.mjs.")


if __name__ == "__main__":
    main()
